import json
from pathlib import Path
from datetime import datetime, timezone
from app.models.meeting import Meeting
from app.pipeline.summarizer import SummaryResult
from app.pipeline.transcriber import TranscriptionResult


OUTPUT_MD_TEMPLATE = """# {title}

## 基本信息
- 日期：{date}
- 时长：{duration}
- 原始文件：{filename}

## 会议摘要
{summary}

## 关键讨论点
{key_points}

## 决策记录
{decisions}

## 行动项
{action_items}

## 标签
{tags}

---

> 由 MeetingSum 自动生成
"""


def _fmt_duration(seconds: int | None) -> str:
    if seconds is None:
        return "--:--:--"
    h = seconds // 3600
    m = (seconds % 3600) // 60
    s = seconds % 60
    return f"{h:02d}:{m:02d}:{s:02d}"


def _format_key_points(key_points: list[dict]) -> str:
    if not key_points:
        return "_无_"
    lines = []
    for i, kp in enumerate(key_points, 1):
        importance = kp.get("importance", "")
        badge = {"high": "【重要】", "medium": "", "low": ""}.get(importance, "")
        lines.append(f"{i}. **{kp.get('topic', '议题')}** {badge}")
        lines.append(f"   {kp.get('content', '')}")
    return "\n".join(lines)


def _format_decisions(decisions: list[dict]) -> str:
    if not decisions:
        return "_无_"
    lines = []
    for d in decisions:
        proposer = d.get("proposer") or "未指定"
        lines.append(f"- [ ] {d.get('content', '')} — 提出人：{proposer}")
    return "\n".join(lines)


def _format_action_items(items: list[dict]) -> str:
    if not items:
        return "_无_"
    lines = []
    for item in items:
        assignee = item.get("assignee") or "待指定"
        deadline = item.get("deadline") or "待定"
        lines.append(f"- [ ] {item.get('content', '')} — 负责人：{assignee} — 截止日期：{deadline}")
    return "\n".join(lines)


def generate_markdown(
    meeting: Meeting,
    summary: SummaryResult,
    output_dir: str,
) -> str:
    """Generate a Markdown summary document and return the file path."""
    output_path = Path(output_dir) / f"{meeting.id}_summary.md"

    date_str = meeting.created_at.strftime("%Y-%m-%d") if meeting.created_at else datetime.now(timezone.utc).strftime("%Y-%m-%d")

    content = OUTPUT_MD_TEMPLATE.format(
        title=summary.title or meeting.title or "未命名会议",
        date=date_str,
        duration=_fmt_duration(meeting.duration_seconds),
        filename=Path(meeting.original_file or "").name or "未知",
        summary=summary.summary,
        key_points=_format_key_points(summary.key_points),
        decisions=_format_decisions(summary.decisions),
        action_items=_format_action_items(summary.action_items),
        tags=", ".join(summary.tags) if summary.tags else "_无_",
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(content, encoding="utf-8")
    return str(output_path)


def generate_transcript_markdown(
    meeting: Meeting,
    transcription: TranscriptionResult,
    output_dir: str,
) -> str:
    """Generate a Markdown file of the full transcript with timestamps."""
    output_path = Path(output_dir) / f"{meeting.id}_transcript.md"

    lines = [
        f"# 完整转录文本 — {meeting.title or '未命名会议'}",
        "",
        f"日期：{meeting.created_at.strftime('%Y-%m-%d') if meeting.created_at else ''}",
        f"时长：{_fmt_duration(meeting.duration_seconds)}",
        "",
        "---",
        "",
        transcription.text_with_timestamps,
        "",
        "> 由 MeetingSum 自动生成",
    ]

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text("\n".join(lines), encoding="utf-8")
    return str(output_path)


def generate_docx(
    meeting: Meeting,
    summary: SummaryResult,
    output_dir: str,
) -> str:
    """Generate a Word (.docx) document and return the file path."""
    from docx import Document
    from docx.shared import Pt, Inches

    output_path = Path(output_dir) / f"{meeting.id}_summary.docx"
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    style = doc.styles["Normal"]
    font = style.font
    font.name = "Arial"
    font.size = Pt(11)

    doc.add_heading(summary.title or meeting.title or "未命名会议", level=1)

    doc.add_heading("基本信息", level=2)
    date_str = meeting.created_at.strftime("%Y-%m-%d") if meeting.created_at else ""
    doc.add_paragraph(f"日期：{date_str}")
    doc.add_paragraph(f"时长：{_fmt_duration(meeting.duration_seconds)}")
    doc.add_paragraph(f"原始文件：{Path(meeting.original_file or '').name or '未知'}")

    doc.add_heading("会议摘要", level=2)
    doc.add_paragraph(summary.summary)

    doc.add_heading("关键讨论点", level=2)
    for i, kp in enumerate(summary.key_points, 1):
        importance = kp.get("importance", "")
        prefix = "【重要】" if importance == "high" else ""
        doc.add_paragraph(f"{i}. {prefix}{kp.get('topic', '议题')}", style="List Number")
        doc.add_paragraph(kp.get("content", ""))

    doc.add_heading("决策记录", level=2)
    for d in summary.decisions:
        proposer = d.get("proposer") or "未指定"
        doc.add_paragraph(f"{d.get('content', '')} — 提出人：{proposer}", style="List Bullet")

    doc.add_heading("行动项", level=2)
    for item in summary.action_items:
        assignee = item.get("assignee") or "待指定"
        deadline = item.get("deadline") or "待定"
        doc.add_paragraph(
            f"{item.get('content', '')} — 负责人：{assignee} — 截止日期：{deadline}",
            style="List Bullet",
        )

    if summary.tags:
        doc.add_heading("标签", level=2)
        doc.add_paragraph(", ".join(summary.tags))

    doc.add_paragraph("")
    doc.add_paragraph("由 MeetingSum 自动生成")

    doc.save(str(output_path))
    return str(output_path)


def export_all(
    meeting: Meeting,
    summary: SummaryResult,
    transcription: TranscriptionResult,
    output_dir: str,
) -> dict[str, str]:
    """Generate all output documents. Returns a dict of format -> path."""
    results = {}

    md_path = generate_markdown(meeting, summary, output_dir)
    results["md"] = md_path

    transcript_path = generate_transcript_markdown(meeting, transcription, output_dir)
    results["transcript"] = transcript_path

    try:
        docx_path = generate_docx(meeting, summary, output_dir)
        results["docx"] = docx_path
    except Exception:
        pass

    return results
